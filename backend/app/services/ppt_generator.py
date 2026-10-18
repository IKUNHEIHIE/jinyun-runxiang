"""Editable lesson slides shared by saved and legacy teacher exports."""
import io
from types import SimpleNamespace

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.core.lesson_plans import local_lesson, text_parts, text_pages

_MISSING = object()

def generate_lesson_pptx(topic: str, grade: str = "小学三年级", duration_minutes: int = 45,
                         targets=_MISSING, key_difficult=_MISSING, teaching_process=_MISSING,
                         teacher_script=_MISSING, interactive_questions=_MISSING, homework_proposal=_MISSING,
                         source="本地课程模板", title=None) -> io.BytesIO:
    if not 10 <= duration_minutes <= 240:
        raise ValueError("课堂时长应为10至240分钟")
    draft = local_lesson(SimpleNamespace(topic=topic, grade=grade, duration_minutes=duration_minutes, objectives=None))
    targets = draft["targets"] if targets is _MISSING else targets
    key_difficult = draft["key_difficult_points"] if key_difficult is _MISSING else key_difficult
    teaching_process = draft["teaching_process"] if teaching_process is _MISSING else teaching_process
    teacher_script = draft["teacher_script"] if teacher_script is _MISSING else teacher_script
    interactive_questions = draft["interactive_questions"] if interactive_questions is _MISSING else interactive_questions
    homework_proposal = draft["homework_proposal"] if homework_proposal is _MISSING else homework_proposal
    title = title or topic + " · 课堂教案"
    sections = [("教学目标", "\n\n".join(text_parts(targets))),
                ("教学重难点", "\n\n".join(text_parts(key_difficult)))]
    if isinstance(teaching_process, list):
        for i, step in enumerate(teaching_process):
            body = (str(step.get("step", "")) + "\n\n" + str(step.get("content", ""))) if isinstance(step, dict) else str(step)
            sections.append(("教学过程 · " + str(i+1), body))
    else:
        sections.append(("教学过程", "\n".join(text_parts(teaching_process))))
    sections += [("课堂讲稿", "\n".join(text_parts(teacher_script))),
                 ("互动问题", "\n\n".join(text_parts(interactive_questions))),
                 ("课后作业", "\n".join(text_parts(homework_proposal)))]
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    ink, accent, muted = RGBColor(74,49,49), RGBColor(159,114,114), RGBColor(110,91,91)

    def textbox(slide, text, x, y, width, height, size, color, bold=False, spacing=1.25):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(str(text).split("\n")):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = color
            paragraph.font.bold = bold
            paragraph.line_spacing = spacing
            paragraph.space_after = Pt(0)
        return box

    def new_slide():
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255,249,249)
        slide.notes_slide.notes_text_frame.text = "来源：" + str(source) + "\n方案：" + str(title) + "\n学段：" + str(grade)
        return slide

    for part in text_pages(title, width=52, lines=5):
        cover = new_slide()
        textbox(cover, "津韵润乡 · 非遗课堂", .8, .65, 11.7, .5, 18, accent)
        textbox(cover, part, .8, 1.55, 11.7, 3.2, 32, ink, True, 1.2)
        metadata = grade + " · " + str(duration_minutes) + "分钟\n来源：" + str(source)
        textbox(cover, "\n".join(text_pages(metadata, width=90, lines=8)), .8, 5.0, 11.7, 1.7, 17, muted, spacing=1.2)
    for heading, body in sections:
        for i, page in enumerate(text_pages(body or "此方案未提供该部分内容")):
            slide = new_slide()
            label = heading + ("（续" + str(i+1) + "）" if i else "")
            textbox(slide, label, .8, .55, 11.7, .85, 32, accent, True)
            textbox(slide, page, .8, 1.65, 11.7, 5.1, 18, ink)
    total = len(presentation.slides)
    for i, slide in enumerate(presentation.slides):
        textbox(slide, "津韵润乡 · 课程教案    " + str(i+1) + " / " + str(total), .8, 7.05, 11.7, .25, 10, muted)
    output = io.BytesIO()
    presentation.save(output)
    output.seek(0)
    return output
