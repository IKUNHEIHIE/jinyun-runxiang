"""Saved lesson flows use isolated DBs and mocked model responses, never paid services."""
import io
import json
import re
from types import SimpleNamespace

import httpx
import pytest
from pptx import Presentation

from test_mobile_client import client, auth
from app.core import lesson_plans
from app.models.entities import GrowthEvent, SysConfig, UserGrowth
from app.models.mobile import LessonPlanRecord

BASE = "/api/mobile/teacher/lesson-plans"
BODY = {"grade":"小学三年级", "topic":"杨柳青年画", "duration_minutes":45, "mode":"LOCAL", "request_id":"lesson-retry-0001"}


def create(client, body=None, uid=2):
    return client.post(BASE, headers=auth(uid,"TEACHER"), json=body or BODY)


def configure(client):
    with client.sessions() as db:
        db.add(SysConfig(config_key="llm_settings",config_value=json.dumps({"api_key":"test-placeholder-only", "api_base":"https://model.invalid/v1", "model_name":"test-model"})))
        db.commit()


def test_lesson_retry_is_scoped_to_teacher_and_payload_and_has_no_rewards(client):
    first = create(client).json()["data"]
    assert create(client).json()["data"]["id"] == first["id"]
    for changed in [{"topic":"古琴鉴赏"},{"grade":"大学"},{"duration_minutes":60},{"mode":"AUTO"}]:
        assert create(client, BODY | changed).status_code == 409
    assert create(client, uid=5).json()["data"]["id"] != first["id"]
    assert create(client, BODY | {"request_id":None}).status_code == 200
    assert create(client, BODY | {"request_id":None}).status_code == 200
    with client.sessions() as db:
        assert db.query(LessonPlanRecord).count() == 4
        assert db.query(GrowthEvent).count() == 0
        assert db.get(UserGrowth,3).energy_value == 100


@pytest.mark.parametrize("changed", [{"grade":"  "},{"topic":"\n\t"},{"duration_minutes":9},{"duration_minutes":241},{"mode":"UNKNOWN"},{"request_id":"bad/slash"}])
def test_lesson_input_bounds_reject_before_persisting(client, changed):
    assert create(client,BODY | changed).status_code == 422
    with client.sessions() as db:
        assert db.query(LessonPlanRecord).count() == 0


def test_lesson_history_cursor_survives_new_record_and_never_leaks_another_teacher(client):
    ids = [create(client,BODY | {"request_id":f"lesson-page-{i}"}).json()["data"]["id"] for i in range(5)]
    other = create(client,uid=5).json()["data"]["id"]
    first = client.get(BASE+"/history?limit=2",headers=auth(2,"TEACHER")).json()["data"]
    assert [r["id"] for r in first["items"]] == ids[::-1][:2] and first["total"] == 5
    assert all(r["created_at"].endswith("+00:00") for r in first["items"])
    new = create(client,BODY | {"request_id":"lesson-added-later"}).json()["data"]["id"]
    accumulated = [r["id"] for r in first["items"]]
    cursor = first["next_cursor"]
    while cursor:
        page = client.get(BASE+f"/history?limit=2&before_id={cursor}",headers=auth(2,"TEACHER")).json()["data"]
        accumulated.extend(r["id"] for r in page["items"])
        cursor = page["next_cursor"]
    assert accumulated == ids[::-1] and other not in accumulated and new not in accumulated
    assert client.get(BASE+"/history?limit=0",headers=auth(2,"TEACHER")).status_code == 422
    assert client.get(BASE+"/history",headers=auth()).status_code == 403


def test_local_mode_never_calls_configured_model_and_capabilities_reveal_no_secret(client, monkeypatch):
    configure(client)
    def forbidden(*args,**kwargs):
        raise AssertionError("LOCAL mode must never contact a model")
    monkeypatch.setattr(lesson_plans.httpx,"Client",forbidden)
    caps = client.get(BASE+"/capabilities",headers=auth(2,"TEACHER")).json()["data"]
    assert caps == {"model_enabled":True,"default_mode":"LOCAL","min_duration_minutes":10,"max_duration_minutes":240}
    assert "placeholder" not in json.dumps(caps)
    data = create(client).json()["data"]
    assert data["source"] == "本地课程模板" and data["content"]["generation_method"] == "local_template"
    assert "ppt_download_url" not in data["content"]


def test_local_lesson_activities_follow_grade_band(client):
    expected = {
        "小学三年级": ("请老师或家长帮助", "两三句话"),
        "初中": ("补充一条可查证的资料", "一段文字"),
        "高中": ("补充两条可查证的资料", "标注所依据"),
        "大学": ("两条以上可查证资料", "事实、推断"),
    }
    contents = {}
    for grade, (homework_marker, skill_marker) in expected.items():
        data = create(client, BODY | {"grade": grade, "request_id": None}).json()["data"]
        content = data["content"]
        assert homework_marker in content["homework_proposal"], grade
        assert skill_marker in content["targets"]["skill"], grade
        assert skill_marker in content["teacher_script"], grade
        contents[grade] = content
    assert len({c["homework_proposal"] for c in contents.values()}) == 4
    assert len({tuple(c["interactive_questions"]) for c in contents.values()}) == 4


def test_html_export_uses_grade_aware_local_content(client):
    params = {"topic": "杨柳青年画", "duration_minutes": 30}
    html_url = "/api/teacher/lesson-plans/export-html"
    high = client.get(html_url, params=params | {"grade": "高中"}, headers=auth(2, "TEACHER"))
    primary = client.get(html_url, params=params | {"grade": "小学三年级"}, headers=auth(2, "TEACHER"))
    assert high.status_code == primary.status_code == 200
    assert "30 分钟" in high.text and "30 分钟" in primary.text
    assert "标注所依据的课堂资料" in high.text
    assert "补充两条可查证的资料" in high.text
    assert "两三句话" in primary.text
    assert "补充一条可查证的资料" in primary.text
    for unsupported in ["AI 诊断三维指标均达标", "激光纸样自动排版", "AI多模态审美门诊", "3D 可动模型检视"]:
        assert unsupported not in high.text
    assert "不包含未接入的 3D 模型" in high.text


@pytest.mark.parametrize("kind", ["incomplete","invalid-json","unavailable","invalid-url"])
def test_configured_model_failures_are_saved_with_honest_fallback_source(client, monkeypatch, kind):
    configure(client)
    class FakeClient:
        def __init__(self,**kwargs): pass
        def __enter__(self): return self
        def __exit__(self,*args): pass
        def post(self,*args,**kwargs):
            if kind == "unavailable": raise httpx.ConnectError("offline")
            if kind == "invalid-url": raise httpx.InvalidURL("invalid test URL")
            value = "not-json" if kind == "invalid-json" else '{"targets":{}}'
            return httpx.Response(200,request=httpx.Request("POST","https://model.invalid/v1"),json={"choices":[{"message":{"content":value}}]})
    monkeypatch.setattr(lesson_plans.httpx,"Client",FakeClient)
    data = create(client,BODY | {"mode":"AUTO"}).json()["data"]
    assert data["source"] == "本地课程模板（模型服务未完成）"
    assert data["content"]["model_fallback"] is True
    assert data["content"]["targets"]["aesthetic"]


def test_valid_configured_model_response_is_normalized_and_saved(client, monkeypatch):
    configure(client)
    content = lesson_plans.local_lesson(SimpleNamespace(**BODY,objectives=None))
    content["targets"]["aesthetic"] = "  模型测试教学目标  "
    content["ppt_download_url"] = "https://example.com/fake.pptx"
    class FakeClient:
        def __init__(self,**kwargs): pass
        def __enter__(self): return self
        def __exit__(self,*args): pass
        def post(self,*args,**kwargs):
            return httpx.Response(200,request=httpx.Request("POST","https://model.invalid/v1"),json={"choices":[{"message":{"content":json.dumps(content)}}]})
    monkeypatch.setattr(lesson_plans.httpx,"Client",FakeClient)
    result = create(client,BODY | {"mode":"AUTO"}).json()["data"]
    assert result["source"] == "已配置模型辅助"
    assert result["content"]["targets"]["aesthetic"] == "模型测试教学目标"
    assert "ppt_download_url" not in result["content"]
    assert client.get(BASE+f'/{result["id"]}',headers=auth(2,"TEACHER")).json()["data"] == result


def test_saved_long_lesson_ppt_keeps_complete_text_and_editable_shapes(client):
    data = create(client).json()["data"]
    script = "\n\n".join(f"第{i}段：观察纹样，记录依据。ASCII-{i}-END。"*10 for i in range(24))
    with client.sessions() as db:
        record = db.get(LessonPlanRecord,data["id"])
        record.content_json = record.content_json | {"teacher_script":script}
        db.commit()
    response = client.get(BASE+f'/{data["id"]}/pptx',headers=auth(2,"TEACHER"))
    assert response.status_code == 200
    presentation = Presentation(io.BytesIO(response.content))
    assert len(presentation.slides) > 15
    script_text = "".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame and shape.top > 1_400_000 and shape.top < 2_000_000)
    assert re.sub(r"\s+","",script) in re.sub(r"\s+","",script_text)
    all_text = "".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
    assert "专家组" not in all_text and "指导单位" not in all_text
    for slide in presentation.slides:
        assert all(shape.has_text_frame for shape in slide.shapes)
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= presentation.slide_width
            assert shape.top+shape.height <= presentation.slide_height


def test_saved_json_keys_do_not_change_the_teaching_order(client):
    result = create(client).json()["data"]
    with client.sessions() as db:
        record = db.get(LessonPlanRecord,result["id"])
        record.content_json = record.content_json | {"targets":{"skill":"技能目标","emotion":"情感目标","aesthetic":"审美目标"}}
        db.commit()
    data = client.get(BASE+f'/{result["id"]}',headers=auth(2,"TEACHER")).json()["data"]
    assert list(data["content"]["targets"]) == ["aesthetic","skill","emotion"]


def test_lesson_pagination_keeps_closing_punctuation_with_its_cjk_line():
    content = "甲"*36+"。"+"乙"*36+"，"+"丙"*36+"？"+"丁"*36+"！"
    lines = "\n".join(lesson_plans.text_pages(content)).splitlines()
    assert "".join(lines) == content
    assert not any(line.startswith(("。","，","？","！")) for line in lines)


@pytest.mark.parametrize("legacy", [{"teacher_script":"旧版保存讲稿","targets":None,"ppt_download_url":"https://example.com/old.pptx"}, "旧版原始讲稿", ["历史片段"]])
def test_legacy_lesson_read_and_export_preserve_database_without_fake_link(client, legacy):
    with client.sessions() as db:
        record = LessonPlanRecord(teacher_id=2,title="历史方案",grade="高中",topic="陶瓷文化",duration_minutes=45,content_json=legacy,source="历史来源")
        db.add(record); db.commit(); plan_id=record.id
    result = client.get(BASE+f"/{plan_id}",headers=auth(2,"TEACHER")).json()["data"]
    assert "ppt_download_url" not in result["content"] and result["content"]["title"] == "历史方案"
    export = client.get(BASE+f"/{plan_id}/pptx",headers=auth(2,"TEACHER"))
    assert export.status_code == 200
    ppt = Presentation(io.BytesIO(export.content))
    exported_text = "".join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
    assert "此方案未提供该部分内容" in exported_text
    assert "观察陶瓷文化的器形" not in exported_text  # A historical NULL is not a new template request.
    with client.sessions() as db:
        assert db.get(LessonPlanRecord,plan_id).content_json == legacy
